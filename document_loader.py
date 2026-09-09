from __future__ import annotations

from io import BytesIO

MAX_EXTRACT_CHARS = 40_000
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt", ".md"}


def _truncate(text: str) -> str:
    text = text.strip()
    if len(text) <= MAX_EXTRACT_CHARS:
        return text
    return (
        text[:MAX_EXTRACT_CHARS].rstrip()
        + "\n\n[Document truncated — only the first portion was used.]"
    )


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return _truncate("\n\n".join(pages))


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(data))
    parts: list[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return _truncate("\n\n".join(parts))


def _extract_plain_text(data: bytes) -> str:
    return _truncate(data.decode("utf-8", errors="replace"))


def extract_text(filename: str, data: bytes) -> str:
    """Extract text from an uploaded project document."""
    if not data:
        raise ValueError("The uploaded file is empty.")

    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type. Supported: {supported}")

    if suffix == ".pdf":
        text = _extract_pdf(data)
    elif suffix == ".docx":
        text = _extract_docx(data)
    elif suffix == ".doc":
        raise ValueError(
            "Legacy .doc files are not supported. Please save the file as .docx and upload again."
        )
    else:
        text = _extract_plain_text(data)

    if not text.strip():
        raise ValueError("No readable text was found in the document.")

    return text
