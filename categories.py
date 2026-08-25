from __future__ import annotations

import re
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

DATA_PATH = Path(__file__).resolve().parent / "data" / "Value_Categories.xlsx"


def _slugify(text: str) -> str:
    text = text.lower().strip()
    text = text.replace("&", " and ").replace("/", " ")
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def _load_taxonomy(path: Path = DATA_PATH) -> list[dict]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    categories: list[dict] = []
    by_name: dict[str, dict] = {}

    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if index == 1:
            continue
        category_name, subcategory_name, description = row
        if not category_name or not subcategory_name:
            continue

        category_name = str(category_name).strip()
        subcategory_name = str(subcategory_name).strip()
        description = str(description or "").strip()

        category = by_name.get(category_name)
        if category is None:
            category = {
                "id": _slugify(category_name),
                "name": category_name,
                "subcategories": [],
            }
            by_name[category_name] = category
            categories.append(category)

        category["subcategories"].append(
            {
                "id": f"{category['id']}__{_slugify(subcategory_name)}",
                "name": subcategory_name,
                "description": description,
            }
        )

    return categories


@lru_cache(maxsize=1)
def get_categories() -> tuple[dict, ...]:
    return tuple(_load_taxonomy())


CATEGORIES = list(get_categories())
CATEGORY_BY_ID = {category["id"]: category for category in CATEGORIES}
CATEGORY_IDS = [category["id"] for category in CATEGORIES]

SUBCATEGORY_BY_ID: dict[str, dict] = {}
for category in CATEGORIES:
    for subcategory in category["subcategories"]:
        SUBCATEGORY_BY_ID[subcategory["id"]] = {
            **subcategory,
            "category_id": category["id"],
            "category_name": category["name"],
        }

SUBCATEGORY_IDS = list(SUBCATEGORY_BY_ID.keys())

DEFAULT_CATEGORY_ID = "project_management_and_delivery"
DEFAULT_SUBCATEGORY_ID = next(
    (
        item["id"]
        for item in CATEGORY_BY_ID[DEFAULT_CATEGORY_ID]["subcategories"]
        if item["name"] == "Project Management"
    ),
    CATEGORY_BY_ID[DEFAULT_CATEGORY_ID]["subcategories"][0]["id"],
)


def category_list_for_prompt() -> str:
    lines: list[str] = []
    for category in CATEGORIES:
        lines.append(f"## {category['id']}: {category['name']}")
        for subcategory in category["subcategories"]:
            lines.append(
                f"- {subcategory['id']}: {subcategory['name']} — {subcategory['description']}"
            )
    return "\n".join(lines)


def subcategories_for(category_id: str) -> list[dict]:
    category = CATEGORY_BY_ID.get(category_id)
    if not category:
        return []
    return list(category["subcategories"])
